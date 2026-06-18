"""Export results to Excel and CSV."""
import logging
from pathlib import Path
from typing import List
from datetime import datetime

import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.utils import get_column_letter

from src.models import ResearchResult

logger = logging.getLogger("ungc.export")

COLUMNS = [
    "source_category", "company_name", "participant_source_url",
    "company_website", "company_domain",
    "target_person_name", "target_person_title",
    "role_match_type", "role_match_score",
    "linkedin_profile_url", "linkedin_confidence", "linkedin_source_url",
    "confirmed_email", "confirmed_email_source_url",
    "email_pattern_found", "email_pattern_evidence", "potential_email", "potential_email_confidence",
    "mx_valid", "other_public_emails_found",
    "research_sources", "confidence_score", "status", "notes", "last_checked",
]


def results_to_df(results: List[ResearchResult]) -> pd.DataFrame:
    rows = []
    for r in results:
        d = r.model_dump()
        d["last_checked"] = d.get("last_checked") or datetime.utcnow().isoformat()
        rows.append(d)
    df = pd.DataFrame(rows, columns=COLUMNS)
    return df


def _apply_header_style(ws, header_fill_color="1F4E79"):
    fill = PatternFill(start_color=header_fill_color, end_color=header_fill_color, fill_type="solid")
    font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

def _auto_width(ws, max_width=60):
    for col in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=10)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 2, max_width)


def export(results: List[ResearchResult], output_path: str) -> None:
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)

    df = results_to_df(results)

    # CSV
    csv_path = path.with_suffix(".csv")
    df.to_csv(csv_path, index=False)
    logger.info(f"CSV saved: {csv_path}")

    # Excel
    wb = openpyxl.Workbook()

    # --- Results sheet ---
    ws_results = wb.active
    ws_results.title = "Results"
    for row in dataframe_to_rows(df, index=False, header=True):
        ws_results.append(row)
    _apply_header_style(ws_results)
    ws_results.freeze_panes = "A2"
    ws_results.auto_filter.ref = ws_results.dimensions
    _auto_width(ws_results)

    # --- Summary sheet ---
    ws_sum = wb.create_sheet("Summary")
    status_counts = df["status"].value_counts().reset_index()
    status_counts.columns = ["Status", "Count"]
    ws_sum.append(["Status", "Count"])
    for _, row in status_counts.iterrows():
        ws_sum.append([row["Status"], row["Count"]])
    ws_sum.append([])
    ws_sum.append(["Total companies", len(df)])
    ws_sum.append(["With confirmed email", int(df["confirmed_email"].notna().sum())])
    ws_sum.append(["With potential email", int(df["potential_email"].notna().sum())])
    ws_sum.append(["With LinkedIn", int(df["linkedin_profile_url"].notna().sum())])
    ws_sum.append(["Person found", int(df["target_person_name"].notna().sum())])
    ws_sum.append(["Avg confidence", round(df["confidence_score"].mean(), 1)])
    _apply_header_style(ws_sum)
    _auto_width(ws_sum)

    # --- Needs Manual Review ---
    review_df = df[df["status"] == "needs_manual_review"]
    ws_review = wb.create_sheet("Needs Manual Review")
    for row in dataframe_to_rows(review_df, index=False, header=True):
        ws_review.append(row)
    _apply_header_style(ws_review)
    ws_review.freeze_panes = "A2"
    _auto_width(ws_review)

    # --- Email Pattern Evidence ---
    pattern_df = df[df["email_pattern_evidence"].notna()][
        ["company_name", "company_domain", "email_pattern_found", "email_pattern_evidence", "potential_email"]
    ]
    ws_pat = wb.create_sheet("Email Pattern Evidence")
    for row in dataframe_to_rows(pattern_df, index=False, header=True):
        ws_pat.append(row)
    _apply_header_style(ws_pat)
    _auto_width(ws_pat)

    # --- Sources Log ---
    sources_df = df[df["research_sources"].notna()][["company_name", "research_sources"]]
    ws_src = wb.create_sheet("Sources Log")
    for row in dataframe_to_rows(sources_df, index=False, header=True):
        ws_src.append(row)
    _apply_header_style(ws_src)
    _auto_width(ws_src)

    wb.save(str(path))
    logger.info(f"Excel saved: {path}")
